from app.models.project import Project
from app.models.task import Task
from app.models.milestone import Milestone
from app.models.raci import RACIAssignment
from app.database import db
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


class ExportService:
    """Сервис для экспорта проектов"""
    
    @staticmethod
    def export_project_to_excel(project_id, user_id):
        """Экспортировать проект в Excel"""
        from app.services.project_service import ProjectService
        
        project, error = ProjectService.get_project(project_id, user_id)
        if error:
            return None, error
        
        wb = openpyxl.Workbook()
        
        ws_info = wb.active
        ws_info.title = "Project Info"
        
        ws_info['A1'] = "Project Report"
        ws_info['A1'].font = Font(size=16, bold=True)
        ws_info['A1'].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws_info['A1'].font = Font(size=16, bold=True, color="FFFFFF")
        ws_info.merge_cells('A1:D1')
        
        info_data = [
            ['Project Name:', project.name],
            ['Description:', project.description or 'N/A'],
            ['Status:', project.status.value],
            ['Priority:', ['Low', 'Medium', 'High'][project.priority]],
            ['Created:', project.created_at.strftime('%Y-%m-%d %H:%M')],
            ['Deadline:', project.deadline.strftime('%Y-%m-%d') if project.deadline else 'N/A'],
            ['Creator:', f"{project.creator.first_name} {project.creator.last_name}" if project.creator else 'N/A'],
        ]
        
        row = 3
        for label, value in info_data:
            ws_info[f'A{row}'] = label
            ws_info[f'A{row}'].font = Font(bold=True)
            ws_info[f'B{row}'] = value
            row += 1
        
        ws_info.column_dimensions['A'].width = 20
        ws_info.column_dimensions['B'].width = 50
        
        ws_milestones = wb.create_sheet("Milestones")
        
        milestone_headers = ['#', 'Name', 'Status', 'Tasks', 'Completed', 'Progress', 'Deadline']
        ws_milestones.append(milestone_headers)
        
        for cell in ws_milestones[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        milestones = Milestone.query.filter_by(project_id=project_id).order_by(Milestone.order).all()
        
        for milestone in milestones:
            milestone_tasks = Task.query.filter_by(milestone_id=milestone.id).all()
            total = len(milestone_tasks)
            completed = sum(1 for t in milestone_tasks if t.status.value == 'DONE')
            progress = f"{(completed/total*100):.0f}%" if total > 0 else "0%"
            
            ws_milestones.append([
                milestone.order,
                milestone.name,
                milestone.status.value,
                total,
                completed,
                progress,
                milestone.deadline.strftime('%Y-%m-%d') if milestone.deadline else 'N/A'
            ])
        
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            ws_milestones.column_dimensions[col].width = 15
        ws_milestones.column_dimensions['B'].width = 30
        
        ws_tasks = wb.create_sheet("Tasks")
        
        task_headers = ['ID', 'Title', 'Milestone', 'Status', 'Priority', 'Assignee (R)', 'Deadline']
        ws_tasks.append(task_headers)
        
        for cell in ws_tasks[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        tasks = Task.query.filter_by(project_id=project_id).order_by(Task.created_at.desc()).all()
        
        for task in tasks:
            responsible = RACIAssignment.query.filter_by(
                task_id=task.id,
                role='RESPONSIBLE'
            ).first()
            
            assignee_name = 'Unassigned'
            if responsible and responsible.user:
                assignee_name = f"{responsible.user.first_name} {responsible.user.last_name}"
            
            milestone_name = task.milestone.name if task.milestone else 'No Milestone'
            priority_text = ['Low', 'Medium', 'High'][task.priority]
            
            ws_tasks.append([
                task.id,
                task.title,
                milestone_name,
                task.status.value,
                priority_text,
                assignee_name,
                task.deadline.strftime('%Y-%m-%d') if task.deadline else 'N/A'
            ])
        
        ws_tasks.column_dimensions['A'].width = 8
        ws_tasks.column_dimensions['B'].width = 40
        ws_tasks.column_dimensions['C'].width = 20
        ws_tasks.column_dimensions['D'].width = 15
        ws_tasks.column_dimensions['E'].width = 12
        ws_tasks.column_dimensions['F'].width = 25
        ws_tasks.column_dimensions['G'].width = 12
        
        ws_raci = wb.create_sheet("RACI Matrix")
        
        team_query = db.session.query(RACIAssignment.user_id).filter(
            RACIAssignment.task_id.in_([t.id for t in tasks])
        ).distinct()
        
        team_user_ids = [u[0] for u in team_query.all()]
        
        from app.models.user import User
        team_members = User.query.filter(User.id.in_(team_user_ids)).all() if team_user_ids else []
        
        raci_headers = ['Task'] + [f"{u.first_name} {u.last_name}" for u in team_members]
        ws_raci.append(raci_headers)
        
        for cell in ws_raci[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
        
        for task in tasks:
            row_data = [task.title]
            
            for member in team_members:
                assignment = RACIAssignment.query.filter_by(
                    task_id=task.id,
                    user_id=member.id
                ).first()
                
                role_text = assignment.role.value[0] if assignment else ''  
                row_data.append(role_text)
            
            ws_raci.append(row_data)
        
        ws_raci.column_dimensions['A'].width = 40
        for i in range(len(team_members)):
            col_letter = chr(66 + i)  
            ws_raci.column_dimensions[col_letter].width = 15
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"project_{project.id}_{project.name.replace(' ', '_')}.xlsx"
        
        return output, filename, None
    
    @staticmethod
    def export_project_to_pdf(project_id, user_id):
        """Экспортировать проект в PDF"""
        from app.services.project_service import ProjectService
        
        project, error = ProjectService.get_project(project_id, user_id)
        if error:
            return None, None, error
        
        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4))
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#4472C4'),
            spaceAfter=30
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#4472C4'),
            spaceAfter=12
        )
        
        elements = []
        
        elements.append(Paragraph(f"Project Report: {project.name}", title_style))
        elements.append(Spacer(1, 0.2*inch))
        
        info_data = [
            ['Project Name:', project.name],
            ['Description:', project.description or 'N/A'],
            ['Status:', project.status.value],
            ['Priority:', ['Low', 'Medium', 'High'][project.priority]],
            ['Created:', project.created_at.strftime('%Y-%m-%d %H:%M')],
            ['Deadline:', project.deadline.strftime('%Y-%m-%d') if project.deadline else 'N/A'],
        ]
        
        info_table = Table(info_data, colWidths=[2*inch, 6*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#E7E6E6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        
        elements.append(info_table)
        elements.append(Spacer(1, 0.3*inch))
        
        elements.append(Paragraph("Milestones", heading_style))
        
        milestones = Milestone.query.filter_by(project_id=project_id).order_by(Milestone.order).all()
        
        if milestones:
            milestone_data = [['#', 'Name', 'Status', 'Tasks', 'Progress', 'Deadline']]
            
            for milestone in milestones:
                milestone_tasks = Task.query.filter_by(milestone_id=milestone.id).all()
                total = len(milestone_tasks)
                completed = sum(1 for t in milestone_tasks if t.status.value == 'DONE')
                progress = f"{(completed/total*100):.0f}%" if total > 0 else "0%"
                
                milestone_data.append([
                    str(milestone.order),
                    milestone.name,
                    milestone.status.value,
                    str(total),
                    progress,
                    milestone.deadline.strftime('%Y-%m-%d') if milestone.deadline else 'N/A'
                ])
            
            milestone_table = Table(milestone_data, colWidths=[0.5*inch, 3*inch, 1.5*inch, 1*inch, 1*inch, 1.5*inch])
            milestone_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            elements.append(milestone_table)
        else:
            elements.append(Paragraph("No milestones", styles['Normal']))
        
        elements.append(PageBreak())
        
        elements.append(Paragraph("Tasks", heading_style))
        
        tasks = Task.query.filter_by(project_id=project_id).order_by(Task.created_at.desc()).all()
        
        if tasks:
            task_data = [['ID', 'Title', 'Status', 'Priority', 'Deadline']]
            
            for task in tasks[:20]: 
                priority_text = ['Low', 'Med', 'High'][task.priority]
                
                task_data.append([
                    str(task.id),
                    task.title[:50],
                    task.status.value,
                    priority_text,
                    task.deadline.strftime('%Y-%m-%d') if task.deadline else 'N/A'
                ])
            
            task_table = Table(task_data, colWidths=[0.5*inch, 5*inch, 1.5*inch, 1*inch, 1.5*inch])
            task_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            elements.append(task_table)
            
            if len(tasks) > 20:
                elements.append(Spacer(1, 0.2*inch))
                elements.append(Paragraph(f"... and {len(tasks) - 20} more tasks", styles['Italic']))
        else:
            elements.append(Paragraph("No tasks", styles['Normal']))
        
        doc.build(elements)
        output.seek(0)
        
        filename = f"project_{project.id}_{project.name.replace(' ', '_')}.pdf"
        
        return output, filename, None
